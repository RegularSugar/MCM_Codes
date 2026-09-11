from openpyxl import load_workbook
p=r'D:\AAA数模\CUMCM2026Problems\C题\附件\附件5\result2.xlsx'
w=load_workbook(p)
for s in w.worksheets:
 print('SHEET',s.title)
 for row in s.iter_rows():
  vals=[(c.coordinate,c.value) for c in row if c.value not in (None,'')]
  if vals: print(vals[:12])
  if row[0].row>8: break
