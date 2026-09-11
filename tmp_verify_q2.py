from openpyxl import load_workbook
import csv
from datetime import datetime
p=r'D:\AAA数模\参赛材料\第二问\outputs\result2.xlsx'
w=load_workbook(p,read_only=True,data_only=False)
for i,s in enumerate(w.worksheets): print(i,s.max_row,s.max_column,repr(s.cell(1,1).value),s.cell(2,1).value)
plan,batt,em=w.worksheets
assert plan.max_row==335 and plan.max_column>=145
assert batt.max_row==1+334*6
assert em.max_row==1+2181
assert em.cell(2,1).value.date()>=datetime(2025,2,1).date()
for target in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']:
    found=[]
    for r in range(2,em.max_row+1):
        d=em.cell(r,1).value
        if d and d.date().isoformat()==target: found.append((em.cell(r,2).value,em.cell(r,3).value))
    print(target,found)
print('q2 workbook validation OK')
