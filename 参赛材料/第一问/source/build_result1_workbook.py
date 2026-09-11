"""将第一问 MILP 的 CSV 结果写入 result1.xlsx 模板。"""
from __future__ import annotations

import csv
import shutil
from pathlib import Path

from openpyxl import load_workbook


def locate_project_dir() -> Path:
    """Locate the project root whether this script is run from code/ or source/."""
    script_dir = Path(__file__).resolve().parent
    for candidate in (script_dir, *script_dir.parents):
        attachment = candidate / "CUMCM2026Problems" / "C题" / "附件" / "附件1.xlsx"
        if attachment.is_file():
            return candidate
    raise FileNotFoundError("未找到附件1.xlsx。")


def locate_output_dir(project_dir: Path) -> Path:
    script_dir = Path(__file__).resolve().parent
    if script_dir.name == "source" and script_dir.parent.name == "第一问":
        return script_dir.parent / "outputs"
    return project_dir / "outputs" / "first_question"


def read_column(path: Path, column: str) -> list[float]:
    with path.open(encoding="utf-8", newline="") as stream:
        return [float(row[column]) for row in csv.DictReader(stream)]


def main() -> None:
    project_dir = locate_project_dir()
    output_dir = locate_output_dir(project_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    data_dir = output_dir if Path(__file__).resolve().parent.name == "source" else project_dir / "code" / "outputs"
    template_path = project_dir / "CUMCM2026Problems" / "C题" / "附件" / "附件5" / "result1.xlsx"
    result_path = output_dir / "result1_第一问.xlsx"

    purchases = read_column(data_dir / "q1_schedule_template_order.csv", "grid_purchase_kWh")
    charges = read_column(data_dir / "q1_battery_summary.csv", "charge_kWh")
    discharges = read_column(data_dir / "q1_battery_summary.csv", "discharge_kWh")
    if len(purchases) != 144 or len(charges) != 6 or len(discharges) != 6:
        raise ValueError("请先运行 solve_q1_milp.py。")

    
    shutil.copy2(template_path, result_path)
    workbook = load_workbook(result_path)
    if workbook.sheetnames != ["计划购电量", "充放电量"]:
        raise ValueError("模板格式错误。")
    plan_sheet = workbook["计划购电量"]
    battery_sheet = workbook["充放电量"]
    original_labels = [plan_sheet.cell(row=row, column=1).value for row in range(2, 146)]

    for row, value in enumerate(purchases, start=2):
        plan_sheet.cell(row=row, column=2).value = round(value, 4)
        plan_sheet.cell(row=row, column=2).number_format = "0.0000"
    for row, (charge, discharge) in enumerate(zip(charges, discharges), start=2):
        battery_sheet.cell(row=row, column=2).value = round(charge, 4)
        battery_sheet.cell(row=row, column=3).value = round(discharge, 4)
        battery_sheet.cell(row=row, column=2).number_format = "0.0000"
        battery_sheet.cell(row=row, column=3).number_format = "0.0000"
    for row in (2, 3):
        battery_sheet.cell(row=row, column=5).value = 6000.0
        battery_sheet.cell(row=row, column=5).number_format = "0.0000"

    final_labels = [plan_sheet.cell(row=row, column=1).value for row in range(2, 146)]
    if final_labels != original_labels:
        raise RuntimeError("模板标签错误。")
    workbook.save(result_path)
    print(f"Saved: {result_path}")


if __name__ == "__main__":
    main()
